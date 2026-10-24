"""在隔离进程中逐页读取PDF与PPTX，保留页码、文字和图片。"""
import json
from pathlib import Path
import sys
import zipfile

MAX_PAGES = 120


def read_pdf(source, folder):
    import pymupdf
    pages = []
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError('PDF需要密码，请先导出可打开的副本。')
        if not 0 < len(document) <= MAX_PAGES:
            raise ValueError('每份课件最多120页，请分成章节上传。')
        for index, page in enumerate(document, 1):
            path = folder / f'page-{index:03}.jpg'
            scale = min(1.8, 1600 / max(page.rect.width, page.rect.height))
            page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(path)
            pages.append({'page': index, 'text': page.get_text('text', sort=True),
                          'images': [str(path)], 'kind': 'pdf_page'})
    return pages


def read_pptx(source, folder):
    from pptx import Presentation
    with zipfile.ZipFile(source) as package:
        if sum(item.file_size for item in package.infolist()) > 300 * 1024 * 1024:
            raise ValueError('课件展开后过大，请按章节拆分。')
    slides = Presentation(source).slides
    if not 0 < len(slides) <= MAX_PAGES:
        raise ValueError('每份课件最多120页，请分成章节上传。')
    pages = []
    for index, slide in enumerate(slides, 1):
        texts, images, shapes = [], [], []
        def visit(items):
            for shape in items:
                if hasattr(shape, 'shapes'):
                    visit(shape.shapes)
                if getattr(shape, 'has_text_frame', False):
                    texts.append(shape.text)
                    shapes.append({'text': shape.text, 'x': int(shape.left), 'y': int(shape.top)})
                if getattr(shape, 'has_table', False):
                    texts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                if hasattr(shape, 'image'):
                    image = shape.image
                    if image.ext.lower() in {'png', 'jpg', 'jpeg', 'gif', 'webp', 'bmp', 'tiff'}:
                        from PIL import Image
                        import io
                        path = folder / f'slide-{index:03}-image-{len(images)+1:02}.jpg'
                        with Image.open(io.BytesIO(image.blob)) as pic:
                            pic.convert('RGB').save(path, quality=90)
                        images.append(str(path))
                if getattr(shape, 'has_chart', False):
                    for series in shape.chart.series:
                        texts.append(str(series.name) + ': ' + ', '.join(map(str, series.values)))
        visit(slide.shapes)
        # 数学对象保留原XML，避免分式、上下标变成缺少结构的字符拼接。
        equations = slide.element.xpath('.//*[local-name()="oMath"]')
        from lxml import etree
        pages.append({'page': index, 'text': '\n'.join(texts), 'images': images,
                      'text_positions': shapes, 'kind': 'pptx_slide',
                      'equations': [etree.tostring(x, encoding='unicode') for x in equations]})
    return pages


def main():
    source, folder, output = map(Path, sys.argv[1:4])
    folder.mkdir(parents=True, exist_ok=True)
    try:
        pages = read_pdf(source, folder) if source.suffix.lower() == '.pdf' else read_pptx(source, folder)
        result = {'ok': True, 'pages': pages}
    except Exception as exc:
        result = {'ok': False, 'message': str(exc) if isinstance(exc, ValueError)
                  else '文件未能读取，请检查文件，或将课件导出为PDF。'}
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')


if __name__ == '__main__':
    main()
